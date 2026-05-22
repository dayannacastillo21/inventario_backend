# -*- coding: utf-8 -*-
"""Genera Guia_Auditoria_APF02_Cafedronel.docx con la auditoria completa APF02."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Guia_Auditoria_APF02_Cafedronel.docx"


def setup_doc():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Cm(2.5)
    s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(2.8)
    s.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    return doc


def h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(doc, text="", bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)


def code(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    r.font.size = Pt(9)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = head
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()
    return t


def main():
    doc = setup_doc()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Guía de Auditoría y Evidencias APF02")
    tr.bold = True
    tr.font.size = Pt(18)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Backend Cafedronel — inventario_backend").italic = True
    para(doc, "Documento generado a partir del código real del repositorio. "
              "Repositorio: https://github.com/dayannacastillo21/inventario_backend")
    doc.add_page_break()

    # 1
    h(doc, "1. Diagnóstico general del proyecto", 1)
    para(doc, "Conclusión: Sí puedes demostrar nivel Excelente si levantas la app, pruebas en Postman "
              "y guardas capturas. El código cumple la rúbrica; falta principalmente evidencia visual.")
    table(doc, ["Ítem", "Valor detectado en el proyecto"], [
        ["Framework", "Spring Boot 3.5.8, Java 17"],
        ["Build", "Maven (pom.xml)"],
        ["Puerto", "8081 (server.port=${PORT:8081})"],
        ["Base de datos", "PostgreSQL en ejecución"],
        ["Tests", "H2 en memoria"],
        ["Config principal", "src/main/resources/application.properties"],
        ["Perfil Spring", "Solo default (no hay spring.profiles.active)"],
        ["Esquema BD", "Flyway V0–V7"],
        ["Swagger", "No existe"],
        ["Pruebas automáticas", "46 tests (mvn test)"],
    ])
    h(doc, "Comandos de ejecución", 2)
    code(doc, "mvn test\nmvn -DskipTests package\nmvn spring-boot:run")
    para(doc, "Variables obligatorias: DB_URL, DB_USER, DB_PASSWORD, JWT_SECRET")
    code(doc, '$env:DB_URL="jdbc:postgresql://localhost:5432/inventario_backenddatabase"\n'
              '$env:DB_USER="postgres"\n$env:DB_PASSWORD="<tu-password>"\n'
              '$env:JWT_SECRET="<clave-minimo-32-bytes>"\n'
              "mvn spring-boot:run")
    table(doc, ["Recurso", "URL"], [
        ["API base", "http://localhost:8081"],
        ["Healthcheck", "http://localhost:8081/api/estado"],
        ["Swagger", "No aplica"],
        ["Postman", "docs/postman/Cafedronel-APF02.postman_collection.json"],
    ])
    doc.add_page_break()

    # 2
    h(doc, "2. Mapeo real del código", 1)
    table(doc,
          ["Entidad", "Tabla", "Repository", "Service", "Controller", "Endpoints", "Relaciones", "Estado"],
          [
              ["Usuario", "usuarios", "UsuarioRepository", "UsuarioServiceImpl",
               "UsuarioController, AuthController", "/api/usuarios, /api/auth/sesiones", "—", "Completo"],
              ["Producto", "productos", "ProductoRepository", "ProductoServiceImpl",
               "ProductoController", "/api/productos", "—", "Completo"],
              ["Proveedor", "proveedores", "ProveedorRepository", "ProveedorServiceImpl",
               "ProveedorController", "/api/proveedores", "—", "Completo"],
              ["Inventario", "inventario", "InventarioRepository", "InventarioServiceImpl",
               "InventarioController", "/api/inventario", "ManyToOne Proveedor", "Completo"],
              ["Pedido", "pedidos", "PedidoRepository", "PedidoServiceImpl",
               "PedidoController", "/api/pedidos", "OneToMany DetallePedido", "Completo"],
              ["DetallePedido", "detalle_pedido", "DetallePedidoRepository", "DetallePedidoServiceImpl",
               "Sin controller propio", "GET /api/pedidos/{id}/detalles", "ManyToOne Pedido, Producto", "Parcial"],
              ["Venta", "ventas", "VentaRepository", "VentaServiceImpl",
               "VentaController", "/api/ventas", "ManyToOne Usuario, Producto", "Completo"],
          ])
    doc.add_page_break()

    # 3 Postman - condensed but complete
    h(doc, "3. Pruebas para Postman", 1)
    para(doc, "Configuración: baseUrl = http://localhost:8081. Login primero. "
              "Authorization > Bearer Token con el campo token de la respuesta.")
    h(doc, "Login (obligatorio primero)", 2)
    table(doc, ["Campo", "Valor"], [
        ["Método", "POST"],
        ["URL", "{{baseUrl}}/api/auth/sesiones"],
        ["Headers", "Content-Type: application/json"],
        ["Body", '{"email":"admin@cafedronel.com","password":"password"}'],
        ["Esperado", "200 + token, role, userId"],
    ])
    h(doc, "Productos /api/productos", 2)
    table(doc, ["Operación", "Método", "URL", "Body / notas", "HTTP", "Captura"], [
        ["Crear", "POST", "/api/productos", '{"nombre":"Mocaccino","precio":14.5,"categoria":"bebidas","descripcion":"Prueba"}', "201", "id en respuesta"],
        ["Listar", "GET", "/api/productos", "Bearer token", "200", "Array JSON"],
        ["Por ID", "GET", "/api/productos/{id}", "Bearer token", "200", "Detalle"],
        ["Actualizar", "PUT", "/api/productos/{id}", "Mismo JSON precio 16.0", "200", "Cambio visible"],
        ["Eliminar", "DELETE", "/api/productos/{id}", "Token ADMIN", "204", "Status 204"],
        ["Activos", "GET", "/api/productos/activos", "Consulta derivada", "200", "Evidencia consultas"],
        ["Precio min", "GET", "/api/productos/busqueda/precio-minimo?min=8", "JPQL", "200", "Evidencia JPQL"],
    ])
    h(doc, "Proveedores /api/proveedores", 2)
    code(doc, 'POST body: {"nombre":"Proveedor Demo","telefono":"999888777",'
              '"direccion":"Av. Test 100","email":"demo@test.com","activo":true}')
    h(doc, "Inventario /api/inventario", 2)
    para(doc, "Campo proveedor = nombre exacto del proveedor en BD (ej. Distribuidora Café Peru del seed V7).")
    code(doc, 'POST: {"nombreInsumo":"Azucar demo","cantidad":20,"unidad":"kg",'
              '"stockMinimo":5,"precioUnitario":3.5,"proveedor":"Distribuidora Café Peru"}')
    bullet(doc, "POST /api/inventario/{id}/deducciones body: {\"unidades\":5}")
    bullet(doc, "GET /api/inventario/alertas/stock-bajo")
    h(doc, "Pedidos /api/pedidos", 2)
    code(doc, 'POST: {"cliente":"Cliente Demo","detalles":[{"productoId":1,"cantidad":2}]}')
    bullet(doc, "GET /api/pedidos/{id}/detalles")
    bullet(doc, "PUT /api/pedidos/{id}/estado?estado=completado")
    h(doc, "Ventas /api/ventas", 2)
    code(doc, 'POST: {"usuarioId":1,"cantidad":2,"productoId":1,"estado":"completado","metodoPago":"efectivo"}')
    bullet(doc, "GET /api/ventas/usuario/1 y GET /api/ventas/estado/pendiente")
    h(doc, "Usuarios /api/usuarios", 2)
    bullet(doc, "POST registro público sin token; GET/PUT/DELETE requieren ADMIN.")
    doc.add_page_break()

    # 4 JWT
    h(doc, "4. Prueba de login y JWT", 1)
    table(doc, ["Ítem", "Valor real"], [
        ["Endpoint", "POST /api/auth/sesiones"],
        ["Clase", "AuthController + UsuarioServiceImpl + JwtService"],
        ["Usuario seed", "admin@cafedronel.com / password (V7__indices_y_datos_demo.sql)"],
        ["Token en respuesta", "Campo token en LoginResponse JSON"],
        ["Uso Postman", "Authorization > Bearer Token"],
    ])
    h(doc, "Tres pruebas obligatorias", 2)
    table(doc, ["#", "Prueba", "URL", "Config", "Esperado"], [
        ["1", "Sin token", "GET /api/productos", "Sin Authorization", "401 JSON"],
        ["2", "Token inválido", "GET /api/productos", "Bearer token.falso", "401"],
        ["3", "Token válido", "GET /api/productos", "Bearer token login", "200"],
    ])
    doc.add_page_break()

    # 5 Security
    h(doc, "5. Pruebas de Spring Security", 1)
    table(doc, ["Componente", "Archivo"], [
        ["SecurityConfig", "security/SecurityConfig.java"],
        ["Filtro JWT", "security/JwtAuthenticationFilter.java"],
        ["JWT", "security/JwtService.java"],
        ["401", "security/RestAuthenticationEntryPoint.java"],
        ["403", "security/RestAccessDeniedHandler.java"],
        ["PasswordEncoder", "BCryptPasswordEncoder bean"],
    ])
    para(doc, "Nivel rúbrica: Excelente. UserDetailsService es stub; login real vía UsuarioServiceImpl.")
    table(doc, ["Ruta", "Token", "Rol", "Sin token", "Con token"], [
        ["GET /api/estado", "No", "—", "200", "200"],
        ["POST /api/auth/sesiones", "No", "—", "—", "200 login"],
        ["GET /api/productos", "Sí", "USER/ADMIN", "401", "200"],
        ["DELETE /api/productos/{id}", "Sí", "ADMIN", "401", "USER:403 ADMIN:204"],
        ["GET /api/admin/ping", "Sí", "ADMIN", "401", "USER:403 ADMIN:200"],
        ["GET /api/usuarios", "Sí", "ADMIN", "401", "USER:403"],
    ])
    doc.add_page_break()

    # 6 SQL
    h(doc, "6. Consultas SQL (PostgreSQL)", 1)
    sql_block = """
-- USUARIOS
SELECT id, nombre, email, rol, activo FROM usuarios ORDER BY id;
SELECT * FROM usuarios WHERE email = 'admin@cafedronel.com';

-- PRODUCTOS
SELECT * FROM productos ORDER BY id DESC LIMIT 5;

-- INVENTARIO + PROVEEDOR
SELECT i.id, i.nombre_insumo, i.cantidad, p.nombre AS proveedor
FROM inventario i JOIN proveedores p ON p.id = i.proveedor_id;

-- PEDIDOS + DETALLE
SELECT * FROM pedidos ORDER BY id DESC LIMIT 5;
SELECT dp.*, pr.nombre FROM detalle_pedido dp
JOIN productos pr ON pr.id = dp.producto_id WHERE dp.pedido_id = 1;

-- VENTAS
SELECT v.id, u.email, pr.nombre, v.total, v.estado FROM ventas v
JOIN usuarios u ON u.id = v.usuario_id
JOIN productos pr ON pr.id = v.producto_id ORDER BY v.id DESC LIMIT 5;

-- FLYWAY
SELECT version, description, success FROM flyway_schema_history ORDER BY installed_rank;
"""
    code(doc, sql_block.strip())
    doc.add_page_break()

    # 7 JPA
    h(doc, "7. JPA / Hibernate", 1)
    bullet(doc, "7 entidades @Entity; ddl-auto=validate; Flyway V0–V7; open-in-view=false")
    bullet(doc, "Relaciones: ManyToOne (Inventario, DetallePedido, Venta); OneToMany (Pedido)")
    bullet(doc, "No hay OneToOne ni ManyToMany")
    para(doc, "Demostrar: Pedido.java + POST pedido + SELECT en pedidos y detalle_pedido.")
    doc.add_page_break()

    # 8 CRUD
    h(doc, "8. CRUD completo", 1)
    table(doc, ["Recurso", "POST", "GET", "GET ID", "PUT", "DELETE", "Estado"], [
        ["Productos", "Sí", "Sí", "Sí", "Sí", "Sí ADMIN", "Completo"],
        ["Proveedores", "Sí", "Sí", "Sí", "Sí", "Sí ADMIN", "Completo"],
        ["Inventario", "Sí", "Sí", "Sí", "Sí", "Sí ADMIN", "Completo"],
        ["Pedidos", "Sí", "Sí", "Sí", "Sí+estado", "Sí ADMIN", "Completo"],
        ["Ventas", "Sí", "Sí", "Sí", "Sí", "Sí ADMIN", "Completo"],
        ["Usuarios", "Sí público", "Sí ADMIN", "Sí ADMIN", "Sí ADMIN", "Sí ADMIN", "Completo"],
        ["DetallePedido", "vía pedido", "vía pedido", "vía pedido", "vía pedido", "cascade", "Parcial"],
    ])
    h(doc, "Orden de demostración Postman", 2)
    for step in [
        "GET /api/estado",
        "POST /api/auth/sesiones → token",
        "CRUD producto completo",
        "Inventario + deducción + stock-bajo",
        "Pedido + detalles + estado",
        "Venta + consulta por usuario",
        "401 sin token; 403 USER en /api/admin/ping",
        "SQL en pgAdmin",
    ]:
        bullet(doc, step)
    doc.add_page_break()

    # 9 Transacciones
    h(doc, "9. Consultas, transacciones y consistencia", 1)
    bullet(doc, "JPQL: buscarConPrecioMinimo, findConStockBajo")
    bullet(doc, "@Transactional en Producto, Usuario, Proveedor, Inventario, Pedido, Venta")
    para(doc, "Prueba consistencia: POST pedido con 2 productos → 2 filas en detalle_pedido; "
              "producto duplicado en detalles → 400 sin guardar pedido válido.")
    doc.add_page_break()

    # 10 Evidencias
    h(doc, "10. Documentación de evidencias", 1)
    code(doc, """evidencias-apf2/
  01-jpa-hibernate/
  02-crud/
  03-consultas-transacciones/
  04-spring-security/
  05-jwt/
  README_APF2.md""")
    para(doc, "En el repo: AVANCE_PROYECTO_FINAL_02.md, CHECKLIST_APF02.md, colección Postman.")
    doc.add_page_break()

    # 11 Checklist
    h(doc, "11. Checklist para sacar 20/20", 1)
    checks = [
        "mvn test → 46 tests BUILD SUCCESS",
        "PostgreSQL + Flyway V0–V7 al arrancar",
        "Variables DB_* y JWT_SECRET configuradas",
        "GET /api/estado → ok: true",
        "Login admin@cafedronel.com / password",
        "Token Bearer en Postman",
        "CRUD productos, proveedores, inventario, pedidos, ventas",
        "401 sin token; 401 token inválido; 403 sin rol ADMIN",
        "SQL en pgAdmin (JOIN pedido-detalle, ventas)",
        "Capturas en evidencias-apf2/",
        "Entregar AVANCE_PROYECTO_FINAL_02.md",
    ]
    for c in checks:
        bullet(doc, "[ ] " + c)
    doc.add_page_break()

    # 12 Calificación
    h(doc, "12. Calificación estimada según rúbrica", 1)
    table(doc, ["Criterio", "Peso", "Estado", "Nota est."], [
        ["JPA/Hibernate", "20%", "Excelente", "5/5"],
        ["CRUD", "25%", "Excelente", "5/5"],
        ["Consultas y transacciones", "15%", "Excelente", "5/5"],
        ["Spring Security", "20%", "Excelente", "4.5–5/5"],
        ["JWT y documentación", "20%", "Excelente", "4.5–5/5"],
    ])
    para(doc, "Nota probable (código): 19–20/20. Con evidencias en exposición: 20/20.", bold=True)
    doc.add_page_break()

    # 13 Guión
    h(doc, "13. Guión de exposición", 1)
    guion = [
        "Mostrar entidades JPA (Pedido con OneToMany y cascade).",
        "Explicar Flyway + validate: el esquema no lo crea Hibernate.",
        "Demo Postman: login → CRUD producto → pedido con detalles.",
        "Consulta JPQL precio mínimo o stock bajo.",
        "GET productos sin token → 401.",
        "Login → Bearer → GET productos 200.",
        "USER intenta /api/admin/ping → 403.",
        "SQL: pedidos + detalle_pedido + ventas.",
        "mvn test 46 pruebas + carpeta evidencias.",
    ]
    for i, line in enumerate(guion, 1):
        para(doc, f"{i}. {line}")
    doc.add_page_break()

    # 14 Conclusión
    h(doc, "14. Resultado final", 1)
    para(doc, "Sí, puedes demostrar nivel Excelente y aspirar a 20/20 si ejecutas las pruebas "
              "de esta guía y guardas las capturas. No se requiere cambiar código para la rúbrica; "
              "solo demostrar lo que ya existe en el repositorio.", bold=True)
    h(doc, "Prioridades antes de exponer", 2)
    bullet(doc, "Alta: levantar app, Postman completo, capturas")
    bullet(doc, "Alta: carpeta evidencias-apf2")
    bullet(doc, "Media: captura mvn test")
    bullet(doc, "Baja: mencionar que Swagger no está implementado")

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
