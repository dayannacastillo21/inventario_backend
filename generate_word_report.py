from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUTPUT_PATH = r"C:\SpringProjectsnew\inventario_backend\Informe_Avance_Proyecto_Final_Cafedronel.docx"


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(3)
section.right_margin = Cm(2.5)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
styles["Normal"].font.size = Pt(11)
for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
    styles[style_name].font.name = "Calibri"
    styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualiza este índice en Word con clic derecho > Actualizar campo."
    fld_char_sep.append(text)
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_sep)
    run._r.append(fld_char_end)


def p(text="", bold=False, italic=False, align=None, style=None):
    para = doc.add_paragraph(style=style)
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    if align is not None:
        para.alignment = align
    return para


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def bullet(text):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(text)
    return para


def num(text):
    para = doc.add_paragraph(style="List Number")
    para.add_run(text)
    return para


def quote(text):
    p(text, style="Intense Quote")


def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                run.bold = True
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


p("INFORME TÉCNICO DEL AVANCE DEL PROYECTO FINAL 01", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, style="Title")
p("Backend Cafedronel", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, style="Subtitle")
doc.add_paragraph()
doc.add_paragraph()
for item in [
    "Curso: Desarrollo de Aplicaciones / Backend con Spring Boot",
    "Proyecto: API REST Backend Cafedronel",
    "Tema del avance: estructura Spring Boot, endpoints REST, capas, pruebas y documentación",
    "Integrantes: [Completar con nombres del equipo]",
    "Docente: [Completar]",
    "Institución: [Completar]",
    "Ciclo / Sección: [Completar]",
    "Fecha: 23 de abril de 2026",
]:
    p(item, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
p(
    "Documento elaborado a partir de la revisión técnica del repositorio del proyecto y alineado a la rúbrica del Avance de Proyecto Final 01.",
    italic=True,
    align=WD_ALIGN_PARAGRAPH.CENTER,
)
doc.add_page_break()

heading("Índice", 1)
add_toc(doc.add_paragraph())
doc.add_page_break()

heading("Resumen Ejecutivo", 1)
for text in [
    "El presente informe documenta el estado actual del proyecto Backend Cafedronel, una API REST desarrollada con Spring Boot orientada a la gestión de productos, pedidos, ventas, inventario, proveedores, usuarios y autenticación básica. El documento se ha estructurado tomando como referencia la rúbrica del Avance de Proyecto Final 01, de modo que cada criterio evaluable quede respaldado con evidencia técnica, descripción funcional y análisis del cumplimiento actual del proyecto.",
    "Durante la revisión del repositorio se identificó un avance importante respecto a versiones previas del sistema. El proyecto ya presenta una arquitectura por capas más clara, rutas REST coherentes bajo el prefijo /api, objetos DTO para validación de entrada, manejo global de errores, ocultamiento de la contraseña en respuestas JSON, pruebas automatizadas básicas y documentación README útil para la ejecución local y la comprensión general del avance.",
    "En términos de madurez académica, el sistema ya supera el nivel de prototipo desordenado y se acerca a un nivel de entrega bastante competitivo para la primera rúbrica. No obstante, aún existen aspectos que pueden pulirse para aspirar al puntaje máximo, principalmente la corrección de problemas de codificación UTF-8 en algunos archivos de documentación y mensajes, el fortalecimiento de la evidencia de ejecución local y el aumento de la cobertura de pruebas automatizadas en endpoints adicionales.",
    "Este informe no solo describe lo ya logrado, sino que además funciona como guía práctica para la elaboración de la versión final en Word. Por ello, incluye explicación técnica, tablas de endpoints, matriz de cumplimiento por rúbrica, lista de evidencias sugeridas, conclusiones y un conjunto de recomendaciones puntuales para maximizar la nota final.",
]:
    p(text)

heading("1. Introducción", 1)
for text in [
    "El desarrollo de aplicaciones backend con Spring Boot constituye una base importante para la construcción de sistemas modernos orientados a servicios. En el caso del proyecto Backend Cafedronel, el objetivo del avance es demostrar la capacidad del equipo para estructurar una API REST funcional, organizada y demostrable, aplicando conceptos fundamentales como la separación por capas, la inyección de dependencias, el diseño de endpoints, la validación de entrada, el manejo de errores y la incorporación de pruebas automatizadas básicas.",
    "El sistema implementado se enfoca en un escenario de gestión interna para una cafetería. A nivel funcional, se contemplan operaciones sobre productos, pedidos, ventas, inventario, proveedores y usuarios, además de un flujo de autenticación básica para inicio de sesión. Aunque el proyecto aún trabaja con datos en memoria y no con una base de datos real, el avance muestra criterios de organización que facilitan su crecimiento posterior.",
    "La importancia de este informe radica en que no solo describe el software desarrollado, sino que también permite evidenciar de manera ordenada cómo el equipo responde a los criterios exigidos por la rúbrica. En otras palabras, el informe sirve como puente entre el trabajo técnico realizado en el código y la evaluación académica formal del avance.",
]:
    p(text)

heading("2. Objetivos", 1)
heading("2.1 Objetivo General", 2)
p("Desarrollar y documentar un backend funcional en Spring Boot para la gestión de operaciones principales de Cafedronel, aplicando arquitectura por capas, endpoints REST, validaciones de entrada, pruebas automatizadas básicas y documentación técnica clara, conforme a los criterios de evaluación del Avance de Proyecto Final 01.")
heading("2.2 Objetivos Específicos", 2)
for item in [
    "Implementar un proyecto Spring Boot correctamente estructurado y configurable para ejecución local.",
    "Exponer endpoints REST organizados bajo convenciones consistentes y con uso apropiado de métodos HTTP.",
    "Separar responsabilidades en controladores, servicios, modelos, DTOs, mapeadores y manejo de excepciones.",
    "Aplicar validaciones a las entradas del usuario para mejorar la robustez del backend.",
    "Incorporar pruebas automatizadas básicas de API y lógica de negocio.",
    "Redactar documentación suficiente para comprender, ejecutar y demostrar el sistema.",
]:
    bullet(item)

heading("3. Descripción General del Proyecto", 1)
for text in [
    "Backend Cafedronel es una API REST construida con Java 17 y Spring Boot 3.5.x. Su propósito es servir como capa backend para un sistema de gestión de cafetería, permitiendo consultar, crear, actualizar y eliminar información relacionada con productos, pedidos, ventas, inventario, proveedores y usuarios. También incluye un endpoint de autenticación para validar credenciales de acceso.",
    "La solución actual mantiene la información en memoria. Esta decisión reduce la complejidad del primer avance y permite concentrarse en la arquitectura, el modelado de endpoints y la validación del comportamiento del sistema. Aunque no se dispone de persistencia real, el proyecto ya adopta decisiones estructurales que facilitan una futura migración a una base de datos relacional y a mecanismos de seguridad más robustos.",
]:
    p(text)
heading("3.1 Módulos Funcionales Implementados", 2)
add_table(
    ["Módulo", "Propósito funcional", "Estado actual"],
    [
        ["Productos", "Registrar, listar, consultar, actualizar y eliminar productos del catálogo", "Implementado"],
        ["Pedidos", "Gestionar pedidos y calcular totales a partir del catálogo de productos", "Implementado"],
        ["Ventas", "Registrar ventas y calcular importes usando el catálogo vigente", "Implementado"],
        ["Inventario", "Gestionar insumos y deducir stock mediante endpoint específico", "Implementado"],
        ["Proveedores", "Administrar proveedores relacionados al abastecimiento", "Implementado"],
        ["Usuarios", "Registrar, listar, actualizar y eliminar usuarios", "Implementado"],
        ["Autenticación", "Validar credenciales en inicio de sesión", "Implementado"],
        ["Estado de API", "Verificar disponibilidad del backend", "Implementado"],
    ],
)

heading("4. Tecnologías Utilizadas", 1)
p("La elección del stack tecnológico responde a los objetivos académicos del avance y a la necesidad de contar con herramientas modernas, ampliamente utilizadas y bien documentadas.")
add_table(
    ["Tecnología", "Uso dentro del proyecto", "Justificación"],
    [
        ["Java 17", "Lenguaje principal del backend", "Versión LTS estable y compatible con Spring Boot 3"],
        ["Spring Boot 3.5.x", "Framework base del proyecto", "Facilita el desarrollo rápido de APIs REST"],
        ["Maven", "Gestión de dependencias y build", "Estandariza compilación, pruebas y empaquetado"],
        ["Jakarta Validation", "Validación de entradas", "Permite aplicar reglas declarativas en DTOs"],
        ["JUnit 5", "Pruebas unitarias", "Framework estándar para testing en Java"],
        ["Spring Boot Test / MockMvc", "Pruebas de integración ligera para endpoints", "Permite validar respuestas HTTP sin cliente externo"],
        ["Docker", "Empaquetado y despliegue", "Facilita portabilidad del servicio"],
        ["Render", "Posible despliegue en la nube", "Permite exponer el backend en un entorno web"],
    ],
)

heading("5. Estructura del Proyecto", 1)
p("Una fortaleza importante del proyecto es la estructura por capas que ya puede observarse en el repositorio. El paquete principal contiene subcarpetas que organizan responsabilidades y hacen que el código sea más mantenible y entendible.")
p("La estructura general observada es la siguiente:")
for line in [
    "src/main/java/com/example/backend_cafedronel/",
    "├── config",
    "├── controller",
    "├── dto",
    "├── exception",
    "├── mapper",
    "├── model",
    "└── service",
]:
    quote(line)
heading("5.1 Descripción de Capas", 2)
add_table(
    ["Capa / paquete", "Descripción", "Archivos representativos"],
    [
        ["config", "Configuración transversal del proyecto", "CorsConfig.java"],
        ["controller", "Exposición de endpoints REST", "ProductoController, PedidoController, AuthController"],
        ["dto", "Objetos de entrada validados para requests", "ProductoRequest, LoginRequest, UsuarioUpdateRequest"],
        ["exception", "Excepciones personalizadas y manejo global de errores", "GlobalExceptionHandler, ResourceNotFoundException"],
        ["mapper", "Transformación entre DTOs y modelos", "PedidoMapper"],
        ["model", "Representación del dominio", "Producto, Pedido, Usuario, Venta"],
        ["service", "Lógica de negocio e integración entre módulos", "PedidoServiceImpl, VentaServiceImpl, UsuarioServiceImpl"],
    ],
)

heading("6. Configuración del Entorno y Ejecución", 1)
for text in [
    "El proyecto está configurado para ejecutarse con Java 17 y Maven 3.9+. En el archivo pom.xml se observan dependencias principales para web, validación y pruebas. Además, application.properties configura el puerto con soporte para una variable de entorno PORT, manteniendo 8081 como valor por defecto para entorno local.",
]:
    p(text)
heading("6.1 Requisitos Previos", 2)
for item in [
    "JDK 17 instalado y configurado en la variable PATH.",
    "Maven 3.9 o superior, o en su defecto el wrapper del proyecto.",
    "IDE recomendado: IntelliJ IDEA, VS Code o Spring Tool Suite.",
    "Cliente HTTP opcional para pruebas manuales: Postman, Thunder Client o navegador.",
]:
    bullet(item)
heading("6.2 Comandos de Ejecución", 2)
p("Para compilar, ejecutar pruebas y levantar el proyecto se recomienda usar la siguiente secuencia:")
quote("mvn -version")
quote("mvn clean test")
quote("mvn spring-boot:run")
heading("6.3 Configuración de Puerto", 2)
p("El archivo application.properties utiliza la siguiente lógica de configuración: el servidor toma el valor de la variable de entorno PORT si existe; de lo contrario, usa 8081. Esta decisión es apropiada porque conserva la facilidad de desarrollo local y, al mismo tiempo, permite despliegues en plataformas como Render.")
heading("6.4 Despliegue Adicional", 2)
p("El proyecto también incorpora un Dockerfile multistage y un archivo render.yaml. Esto no es un requisito explícito de la rúbrica, pero suma valor al mostrar preparación para despliegue y empaquetado. Dichos archivos reflejan interés por la portabilidad del servicio.")
heading("6.5 Evidencia Sugerida", 2)
for item in [
    "Insertar captura de consola mostrando la versión de Maven y Java.",
    "Insertar captura de consola con BUILD SUCCESS tras ejecutar mvn clean test.",
    "Insertar captura de la aplicación levantando en el puerto 8081 con mvn spring-boot:run.",
    "Insertar captura de GET /api/estado respondiendo correctamente.",
]:
    num(item)

heading("7. Implementación de Controladores y Endpoints REST", 1)
p("El proyecto ya cuenta con una colección sólida de controladores REST. La principal mejora respecto a una versión más temprana es la estandarización de rutas bajo el prefijo /api y el uso consistente de códigos HTTP adecuados, tales como 200 para consultas exitosas, 201 para creación y 204 para eliminación.")
heading("7.1 Catálogo General de Endpoints", 2)
add_table(
    ["Método", "Ruta", "Descripción", "Código esperado"],
    [
        ["GET", "/api/estado", "Verifica disponibilidad del backend", "200"],
        ["GET", "/api/productos", "Lista productos", "200"],
        ["GET", "/api/productos/{id}", "Obtiene un producto por id", "200/404"],
        ["GET", "/api/productos/categoria/{categoria}", "Filtra productos por categoría", "200"],
        ["POST", "/api/productos", "Crea un producto", "201/400"],
        ["PUT", "/api/productos/{id}", "Actualiza un producto", "200/404/400"],
        ["DELETE", "/api/productos/{id}", "Elimina un producto", "204/404"],
        ["GET", "/api/pedidos", "Lista pedidos", "200"],
        ["GET", "/api/pedidos/{id}", "Consulta pedido por id", "200/404"],
        ["GET", "/api/pedidos/{id}/detalles", "Lista detalles de un pedido", "200/404"],
        ["POST", "/api/pedidos", "Crea un pedido y calcula total", "201/400/404"],
        ["PUT", "/api/pedidos/{id}", "Actualiza un pedido", "200/400/404"],
        ["PUT", "/api/pedidos/{id}/estado", "Cambia el estado del pedido", "200/400/404"],
        ["DELETE", "/api/pedidos/{id}", "Elimina un pedido", "204/404"],
        ["GET", "/api/ventas", "Lista ventas", "200"],
        ["POST", "/api/ventas", "Registra una venta", "201/400/404"],
        ["PUT", "/api/ventas/{id}", "Actualiza una venta", "200/400/404"],
        ["DELETE", "/api/ventas/{id}", "Elimina una venta", "204/404"],
        ["GET", "/api/inventario", "Lista insumos", "200"],
        ["POST", "/api/inventario", "Crea un insumo", "201/400"],
        ["PUT", "/api/inventario/{id}", "Actualiza un insumo", "200/400/404"],
        ["DELETE", "/api/inventario/{id}", "Elimina un insumo", "204/404"],
        ["POST", "/api/inventario/{id}/deducciones", "Descuenta unidades de stock", "200/400/404"],
        ["GET", "/api/proveedores", "Lista proveedores", "200"],
        ["POST", "/api/proveedores", "Crea un proveedor", "201/400"],
        ["PUT", "/api/proveedores/{id}", "Actualiza un proveedor", "200/400/404"],
        ["DELETE", "/api/proveedores/{id}", "Elimina un proveedor", "204/404"],
        ["GET", "/api/usuarios", "Lista usuarios", "200"],
        ["POST", "/api/usuarios", "Registra usuario", "201/400/409"],
        ["PUT", "/api/usuarios/{id}", "Actualiza usuario", "200/400/404"],
        ["DELETE", "/api/usuarios/{id}", "Elimina usuario", "204/404"],
        ["POST", "/api/auth/sesiones", "Inicia sesión", "200/400/401"],
    ],
)
heading("7.2 Análisis por Módulo", 2)
for title, text in [
    ("7.2.1 Productos", "El módulo de productos representa el catálogo base del sistema. Sus endpoints permiten listar, consultar, filtrar por categoría, crear, actualizar y eliminar. La creación y actualización ya utilizan un DTO específico con validaciones, lo cual mejora la consistencia del contrato de la API."),
    ("7.2.2 Pedidos", "El módulo de pedidos es uno de los puntos más interesantes del proyecto, ya que el total no se recibe directamente desde el cliente, sino que se calcula en el servicio tomando como fuente de verdad el catálogo de productos. Esto demuestra una mejor lógica de negocio y evita inconsistencias derivadas de valores manipulados desde la petición."),
    ("7.2.3 Ventas", "Las ventas se construyen a partir del identificador del producto y vuelven a resolver el precio con el catálogo central. Esta decisión elimina la duplicación de listas locales de productos que existía en versiones anteriores y fortalece la coherencia del backend."),
    ("7.2.4 Inventario", "El inventario permite altas, actualizaciones, bajas y deducción de stock mediante un endpoint dedicado. La deducción valida que el número de unidades sea mayor que cero, lo que evita solicitudes ilógicas."),
    ("7.2.5 Proveedores", "El módulo de proveedores sigue un patrón REST convencional y utiliza un DTO validado para crear y actualizar. La implementación es adecuada para la etapa actual del proyecto."),
    ("7.2.6 Usuarios y autenticación", "El manejo de usuarios incorpora registro, listado, actualización y eliminación. Además, el login fue desacoplado hacia un controlador específico de autenticación. Un detalle importante es que la contraseña ya no se serializa en la respuesta JSON, gracias al uso de JsonProperty con acceso WRITE_ONLY."),
]:
    heading(title, 3)
    p(text)

heading("8. Organización por Capas e Inyección de Dependencias", 1)
for text in [
    "Uno de los criterios centrales de la rúbrica es la separación básica de responsabilidades. El proyecto actual sí evidencia dicha separación. Los controladores se encargan de recibir la solicitud HTTP y delegar la lógica al servicio correspondiente. Los servicios encapsulan reglas del negocio, resolviendo catálogos, calculando subtotales y validando estados antes de responder.",
    "La inyección de dependencias se observa principalmente mediante constructores, lo cual es una práctica recomendada porque mejora la testabilidad y hace explícitas las dependencias de cada componente. Por ejemplo, PedidoServiceImpl recibe ProductoService para resolver los productos desde una fuente central. A su vez, VentaServiceImpl también depende de ProductoService, evitando manejar un catálogo duplicado y mejorando la consistencia del comportamiento.",
    "Otra evidencia de madurez es el uso de un mapper dedicado en la creación y actualización de pedidos. Esto evita que el controlador reciba modelos internos sin control y demuestra que el equipo ya reconoce la diferencia entre el contrato de entrada y la representación interna del dominio.",
]:
    p(text)
heading("8.1 Flujo General de Responsabilidades", 2)
quote("Cliente HTTP → Controller → Service → Model / DTO / Exception → Respuesta JSON")

heading("9. Validaciones y Manejo de Errores", 1)
for text in [
    "La incorporación de DTOs validados constituye uno de los saltos cualitativos más importantes del proyecto. Las solicitudes ya no se reciben de forma totalmente libre sobre los modelos, sino que atraviesan reglas declarativas como @NotBlank, @NotNull, @Positive, @Email y @Size. Esto incrementa la robustez del sistema y evita errores por datos vacíos, negativos o mal formados.",
    "El manejo de errores también ha mejorado gracias a un GlobalExceptionHandler. Esta clase centraliza la conversión de excepciones del dominio a respuestas HTTP bien definidas, permitiendo que la API devuelva mensajes consistentes cuando un recurso no existe, cuando el cliente envía datos inválidos, cuando hay credenciales incorrectas o cuando el email ya se encuentra registrado.",
]:
    p(text)
add_table(
    ["Tipo de problema", "Excepción / origen", "Código HTTP"],
    [
        ["Recurso no encontrado", "ResourceNotFoundException", "404"],
        ["Regla de negocio inválida", "BusinessException", "400"],
        ["Email duplicado", "DuplicateEmailException", "409"],
        ["Credenciales incorrectas", "InvalidCredentialsException", "401"],
        ["Validación de DTO", "MethodArgumentNotValidException", "400"],
        ["Error inesperado", "Exception genérica", "500"],
    ],
)

heading("10. Pruebas Automatizadas y Aplicación Inicial de TDD", 1)
for text in [
    "La rúbrica exige la presencia de pruebas básicas asociadas a la lógica o a los endpoints. En el estado actual del proyecto ya se observan cinco clases de prueba: una para carga del contexto, dos orientadas a API con MockMvc y dos orientadas a lógica de negocio. Este avance permite sostener con mayor seguridad que el sistema no solo fue programado, sino también validado.",
]:
    p(text)
add_table(
    ["Prueba", "Tipo", "Comportamiento validado"],
    [
        ["BackendCafedronelApplicationTests", "Contexto Spring", "Verifica que la aplicación cargue correctamente"],
        ["ProductoApiTest", "MockMvc", "Valida GET 200, GET 404, POST 201 y POST 400"],
        ["AuthApiTest", "MockMvc", "Valida respuesta 401 para login inválido"],
        ["PedidoServiceImplTest", "Unitaria con mock", "Valida cálculo de total y subtotal usando catálogo"],
        ["UsuarioServiceImplTest", "Unitaria", "Valida rechazo de emails duplicados"],
    ],
)
p("Las pruebas actuales son adecuadas para un avance inicial porque cubren tanto comportamiento HTTP como lógica de negocio. Aun así, para alcanzar un 20 completamente sólido sería recomendable añadir algunas pruebas más: creación válida de usuario, actualización de recurso inexistente, validación de inventario y un caso exitoso de autenticación.")
heading("10.1 Evidencia Sugerida para el Informe", 2)
for item in [
    "Insertar captura de mvn clean test finalizando con BUILD SUCCESS.",
    "Insertar captura de clases de prueba visibles en el IDE o repositorio.",
    "Insertar, si es posible, un breve extracto de consola donde se aprecie la ejecución de ProductoApiTest y AuthApiTest.",
]:
    num(item)

heading("11. Documentación Técnica del Avance", 1)
for text in [
    "La documentación mínima solicitada por la rúbrica ya está atendida mediante un README.md bastante más completo que en versiones anteriores. Este documento incluye descripción, objetivo, tecnologías, estructura del proyecto, comandos de ejecución, tabla de endpoints, ejemplos JSON, pruebas automatizadas y evidencias sugeridas.",
    "Desde la perspectiva académica, el README ya cumple una función doble: sirve como guía rápida de ejecución y como resumen técnico del avance. Sin embargo, todavía existe un aspecto de presentación por corregir: algunos textos muestran problemas de codificación de caracteres, por ejemplo palabras con tildes y símbolos especiales. Aunque esto no impide comprender el contenido, sí reduce la percepción de pulcritud final del documento.",
]:
    p(text)
heading("11.1 Recomendación de Mejora Documental", 2)
for item in [
    "Guardar README, comentarios y archivos fuente usando UTF-8.",
    "Revisar texto con tildes, eñes y flechas especiales antes de la entrega final.",
    "Acompañar el README con este informe en Word y capturas de funcionamiento.",
]:
    bullet(item)

heading("12. Análisis del Cumplimiento según la Rúbrica", 1)
p("A continuación se presenta una matriz de cumplimiento alineada directamente a los cinco criterios de evaluación del Avance de Proyecto Final 01.")
add_table(
    ["Criterio de la rúbrica", "Evidencia encontrada en el proyecto", "Nivel de cumplimiento"],
    [
        ["Estructura del proyecto Spring Boot y configuración del entorno", "Proyecto Spring Boot válido, pom.xml organizado, application.properties configurable, Dockerfile y render.yaml presentes, README con ejecución", "Bueno / Muy bueno"],
        ["Implementación de controladores y endpoints REST", "Controladores REST bajo /api, rutas coherentes, códigos HTTP más consistentes, endpoint de estado y autenticación", "Muy bueno"],
        ["Uso de inyección de dependencias y organización por capas", "Separación controller/service/model/dto/exception/mapper y uso de inyección por constructor", "Muy bueno"],
        ["Aplicación inicial de TDD y pruebas del proyecto", "Pruebas unitarias y MockMvc presentes con casos de éxito y error", "Bueno / Muy bueno"],
        ["Documentación técnica básica del avance", "README extenso y este informe técnico estructurado", "Muy bueno"],
    ],
)
heading("12.1 Nota Estimada", 2)
p("Sobre la base de la revisión del código y del alineamiento con la rúbrica, la nota estimada del proyecto se ubica en un rango alto. Si el equipo acompaña la entrega con evidencia de compilación y pruebas exitosas, una defensa razonable sería entre 18.5 y 19 puntos sobre 20. El máximo puntaje es alcanzable si se corrigen los últimos detalles de presentación y se demuestra la ejecución local sin fallos.")

heading("13. Fortalezas del Avance", 1)
for item in [
    "La estructura del proyecto ya es entendible, mantenible y académicamente defendible.",
    "Los endpoints REST siguen convenciones más consistentes que una implementación improvisada.",
    "Los DTOs incorporan validaciones útiles y muestran comprensión de buenas prácticas.",
    "El manejo global de errores mejora la experiencia del cliente de la API.",
    "Las contraseñas dejaron de exponerse en respuestas JSON.",
    "Los servicios de pedidos y ventas resuelven precios desde un catálogo común, evitando incoherencias.",
    "Las pruebas automatizadas ya son reales y no meramente simbólicas.",
    "El README contiene información suficiente para comprender y ejecutar el proyecto.",
]:
    bullet(item)

heading("14. Aspectos Pendientes para Asegurar el 20", 1)
p("Aunque el nivel actual del proyecto es bastante competitivo, todavía existen detalles concretos que conviene cerrar antes de la entrega final. Estos puntos no invalidan el avance, pero sí pueden marcar la diferencia entre una nota muy buena y una sobresaliente.")
add_table(
    ["Pendiente", "Impacto en la evaluación", "Acción recomendada"],
    [
        ["Corregir codificación UTF-8", "Mejora la presentación del README, comentarios y mensajes", "Guardar archivos como UTF-8 y revisar caracteres especiales"],
        ["Mostrar build exitoso", "Fortalece criterio de ejecución verificable", "Adjuntar captura o demo de mvn clean test"],
        ["Mostrar arranque exitoso del backend", "Refuerza entorno y funcionamiento local", "Adjuntar captura de mvn spring-boot:run"],
        ["Agregar 2 o 3 pruebas adicionales", "Sube percepción de madurez en testing", "Cubrir éxito de login, validación de usuarios o errores 404 adicionales"],
        ["Pulir evidencias visuales en Word", "Mejora impacto de la entrega", "Usar capturas limpias y pie de figura numerado"],
    ],
)

heading("15. Dificultades Encontradas y Soluciones Aplicadas", 1)
for text in [
    "A lo largo de la evolución del proyecto se reconocen varias dificultades típicas de una primera construcción backend. Inicialmente, algunas rutas eran inconsistentes y varios módulos manejaban listas locales duplicadas de productos. Esto generaba riesgo de incoherencia entre ventas, pedidos y catálogo.",
    "La solución aplicada consistió en reestructurar controladores hacia un prefijo común /api y hacer que los servicios que necesitan productos resuelvan la información a través de ProductoService. De esta forma, el cálculo de pedidos y ventas ya se basa en una referencia de catálogo más centralizada.",
    "Otro problema inicial fue la ausencia de validaciones y de manejo uniforme de errores. La solución consistió en introducir DTOs, anotar los request bodies con @Valid y centralizar el tratamiento de excepciones mediante GlobalExceptionHandler.",
    "También se detectó que la contraseña del usuario estaba expuesta al serializar respuestas. Este problema fue resuelto utilizando JsonProperty con acceso WRITE_ONLY. Se trata de una mejora importante, aunque todavía se recomienda que en una etapa futura las contraseñas pasen a almacenarse de forma cifrada y se implemente seguridad con Spring Security y JWT.",
]:
    p(text)

heading("16. Conclusiones", 1)
for text in [
    "El proyecto Backend Cafedronel demuestra un avance sólido para la etapa APF1. El backend se encuentra estructurado en Spring Boot, cuenta con múltiples endpoints REST, presenta una separación por capas entendible, aplica validaciones básicas, incorpora manejo global de errores y dispone de pruebas automatizadas iniciales.",
    "Desde la perspectiva de la rúbrica, el trabajo ya cumple de manera fuerte en organización por capas, endpoints REST y documentación técnica, además de haber mejorado de forma importante en pruebas y consistencia de la API. Esto permite afirmar que el avance actual no solo cumple, sino que compite por una calificación alta dentro del rango superior.",
    "Sin embargo, la excelencia completa requiere atención a los últimos detalles: evidencias de ejecución y pruebas, corrección del encoding de caracteres y un pequeño refuerzo en la cobertura de testing. Si el equipo atiende estos puntos antes de la entrega definitiva, las probabilidades de obtener el máximo puntaje aumentan considerablemente.",
]:
    p(text)

heading("17. Trabajo Futuro", 1)
for item in [
    "Migrar de datos en memoria a persistencia real con Spring Data JPA y una base de datos relacional.",
    "Incorporar Spring Security y autenticación basada en JWT.",
    "Cifrar contraseñas y aplicar políticas básicas de seguridad.",
    "Ampliar la cobertura de pruebas unitarias y de integración.",
    "Agregar documentación OpenAPI/Swagger para consumo de la API.",
    "Desplegar una versión pública del backend para demostración remota.",
]:
    bullet(item)

heading("18. Anexos", 1)
heading("18.1 Checklist de Evidencias para Insertar en Word", 2)
for item in [
    "Figura 1. Consola mostrando la versión de Java y Maven.",
    "Figura 2. Ejecución exitosa de mvn clean test.",
    "Figura 3. Arranque del backend en el puerto 8081.",
    "Figura 4. Respuesta exitosa de GET /api/estado.",
    "Figura 5. Respuesta exitosa de GET /api/productos.",
    "Figura 6. Creación válida de producto con POST /api/productos.",
    "Figura 7. Validación 400 en producto inválido.",
    "Figura 8. Login inválido devolviendo 401.",
    "Figura 9. Estructura del proyecto visible en el IDE o explorador de archivos.",
]:
    bullet(item)
heading("18.2 Guía de Presentación en Word", 2)
for item in [
    "Usar numeración automática para títulos y subtítulos.",
    "Actualizar el índice automático antes de exportar el documento final.",
    "Agregar pies de figura en todas las capturas.",
    "Mantener una estética sobria: fondo blanco, texto negro y azul oscuro para títulos.",
    "Revisar ortografía, tildes y formato final antes de entregar.",
]:
    bullet(item)
heading("18.3 Fragmentos de Evidencia Técnica", 2)
p("Comandos principales del proyecto:", bold=True)
quote("mvn clean test")
quote("mvn spring-boot:run")
p("Endpoints de demostración rápida:", bold=True)
quote("GET http://localhost:8081/api/estado")
quote("GET http://localhost:8081/api/productos")
quote("POST http://localhost:8081/api/auth/sesiones")

p(
    "Nota final para edición: este documento ya contiene la estructura, la redacción y el análisis técnico base. Antes de entregar, el equipo debe completar los datos de portada, insertar capturas de evidencia y actualizar el índice automático dentro de Microsoft Word.",
    italic=True,
)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
